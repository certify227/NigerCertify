#!/usr/bin/env python3
"""Génère la proposition technique Word — installation solaire autonome Niamey."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches


OUT = Path("/workspace/dossiers-solaire/Proposition_Installation_Solaire_Niamey.docx")
SCHEMA = Path("/workspace/dossiers-solaire/assets/schema-cablage-solaire-niamey.png")

ORANGE = RGBColor(0xE3, 0x6C, 0x09)
NAVY = RGBColor(0x00, 0x22, 0x4B)
BLUE = RGBColor(0x00, 0x55, 0xD2)
GREEN = RGBColor(0x1B, 0x7A, 0x3D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x44, 0x44, 0x44)


def set_run_font(run, size=11, bold=False, color=NAVY, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="BFBFBF", size="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_para(doc, text, size=11, bold=False, color=NAVY, align="left", space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading_bar(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "00224B")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=WHITE)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def fill_header_row(row, headers, fill="E36C09"):
    for i, h in enumerate(headers):
        cell = row.cells[i]
        shade_cell(cell, fill)
        set_cell_borders(cell, "B85400", "8")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True, color=WHITE)


def fill_row(row, values, zebra=False, bold_last=False):
    for i, v in enumerate(values):
        cell = row.cells[i]
        if zebra:
            shade_cell(cell, "F5F7FA")
        set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i <= 1 else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(v))
        set_run_font(run, size=9, bold=(bold_last and i == len(values) - 1), color=NAVY)


def build():
    doc = Document()

    # Marges
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # En-tête
    add_para(doc, "NIGER CERTIFY", size=18, bold=True, color=ORANGE, align="center", space_after=0)
    add_para(
        doc,
        "Expertise · Énergie solaire · Installations autonomes",
        size=10,
        color=BLUE,
        align="center",
        space_after=2,
    )
    add_para(
        doc,
        "Niamey SONUCI · Tél. +227 94 10 70 74 · contact@nigercertify.com",
        size=9,
        color=GRAY,
        align="center",
        space_after=10,
    )

    # Bandeau titre
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade_cell(c, "E36C09")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROPOSITION TECHNIQUE D’INSTALLATION SOLAIRE AUTONOME")
    set_run_font(r, size=13, bold=True, color=WHITE)
    add_para(doc, "Site isolé — Niamey (République du Niger)", size=11, bold=True, color=NAVY, align="center", space_before=6)
    add_para(doc, "Réf. : NC-SOL-2026-09-001  ·  Date : 18/09/2026  ·  Validité : 30 jours", size=9, color=GRAY, align="center", space_after=12)

    # 1. Contexte
    add_heading_bar(doc, "1. Contexte & objectifs")
    add_para(
        doc,
        "Cette proposition dimensionne une installation photovoltaïque 100 % autonome pour un usage domestique "
        "à Niamey, sans raccordement au réseau électrique. L’objectif est d’assurer environ 3 heures d’autonomie "
        "en soirée pour le multimédia, l’éclairage et le confort, avec une évolution prévue vers l’ajout d’un réfrigérateur.",
        size=10,
        color=GRAY,
        align="justify",
        space_after=8,
    )

    # 2. Bilan de besoin
    add_heading_bar(doc, "2. Bilan de besoin (charges électriques)")
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fill_header_row(table.rows[0], ["Équipement", "Puissance", "Durée", "Énergie / soirée", "Priorité"])
    besoins = [
        ("Télévision (moyenne réelle)", "180 W", "3 h", "540 Wh", "Essentiel"),
        ("Décodeur", "15 W", "3 h", "45 Wh", "Essentiel"),
        ("Kit Starlink Mini", "30 W", "3 h", "90 Wh", "Essentiel"),
        ("3 ampoules LED (15 W)", "45 W", "3 h", "135 Wh", "Essentiel"),
        ("Humidificateur", "150 W", "3 h", "450 Wh", "Confort"),
        ("TOTAL SIMULTANÉ / SOIRÉE", "~ 420–520 W", "3 h", "~ 1 260 Wh", "—"),
    ]
    for i, row_data in enumerate(besoins):
        row = table.add_row()
        fill_row(row, row_data, zebra=(i % 2 == 1), bold_last=False)
        if i == len(besoins) - 1:
            for cell in row.cells:
                shade_cell(cell, "D9E1F2")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
    add_para(
        doc,
        "Avec les pertes de conversion de l’onduleur (~15 %), le prélèvement réel sur la batterie est d’environ "
        "1 500 Wh par soirée. Pic de démarrage futur (frigo) : 1 000 à 1 500 W — d’où l’intérêt de conserver l’onduleur 2 000 W.",
        size=9,
        color=GRAY,
        align="justify",
        space_before=6,
        space_after=10,
    )

    # 3. Matériel existant
    add_heading_bar(doc, "3. Matériel déjà disponible")
    table = doc.add_table(rows=1, cols=3)
    fill_header_row(table.rows[0], ["Élément", "Caractéristiques", "Verdict technique"], fill="0055D2")
    existing = [
        ("4 panneaux solaires", "3 × 250 Wc + 1 × 330 Wc (total 1 080 Wc)", "OK — production largement suffisante"),
        ("Régulateur PWM 30 A", "Existant, ouvert au remplacement / complément", "À conserver pour le panneau 330 Wc seul"),
        ("Onduleur 2 000 W", "À confirmer 12 V / Pur Sinus", "OK pour pics (frigo) — limiter à ~500 W aujourd’hui"),
        ("Batterie (projet)", "200 Ah ou 300 Ah / 12 V", "Privilégier 300 Ah GEL (ou Lithium 200 Ah)"),
    ]
    for i, row_data in enumerate(existing):
        row = table.add_row()
        fill_row(row, row_data, zebra=(i % 2 == 1))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Schema
    add_heading_bar(doc, "4. Schéma de câblage recommandé")
    add_para(
        doc,
        "Architecture bi-régulateur : les 3 panneaux 250 Wc alimentent un nouveau MPPT 60 A ; le panneau 330 Wc "
        "reste sur le PWM 30 A existant. Les deux régulateurs chargent la même batterie 12 V, qui alimente l’onduleur 2 000 W.",
        size=10,
        color=GRAY,
        align="justify",
        space_after=8,
    )
    if SCHEMA.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(SCHEMA), width=Cm(17.5))
    add_para(
        doc,
        "Figure 1 — Schéma unifilaire simplifié (panneaux → régulateurs → batterie → onduleur → charges AC).",
        size=8,
        color=GRAY,
        align="center",
        space_after=10,
    )

    # 5. Tableau achats
    add_heading_bar(doc, "5. Tableau des besoins à acheter (priorisé)")
    add_para(
        doc,
        "Les montants indiqués sont des estimations de marché Niamey / Niger (FCFA TTC indicative). "
        "À confirmer selon marques exactes et disponibilités locales.",
        size=9,
        color=GRAY,
        align="justify",
        space_after=6,
    )

    table = doc.add_table(rows=1, cols=6)
    fill_header_row(
        table.rows[0],
        ["Priorité", "Article à acheter", "Spécification technique", "Qté", "Prix unit. estimé (FCFA)", "Sous-total (FCFA)"],
    )
    achats = [
        ("P1 — Critique", "Régulateur MPPT", "12 V / 60 A (ou 40 A mini), entrée PV haute tension", "1", "85 000 – 150 000", "120 000"),
        ("P1 — Critique", "Batterie stationnaire GEL", "12 V 300 Ah (cycle profond) — ou LiFePO4 12 V 200 Ah", "1", "280 000 – 450 000", "350 000"),
        ("P1 — Critique", "Câble batterie / onduleur", "Cuivre 25 mm² ou 35 mm² (rouge + noir), ≤ 1,5 m", "2 × 1,5 m", "8 000 – 15 000 /m", "35 000"),
        ("P1 — Critique", "Fusible Mega-Fuse + porte-fusible", "125 A DC, sur positif batterie", "1", "12 000 – 25 000", "18 000"),
        ("P2 — Sécurité", "Câble solaire", "6 mm² double isolation UV", "10–15 m", "2 500 – 4 000 /m", "40 000"),
        ("P2 — Sécurité", "Connecteurs MC4 + borniers", "Paires MC4 + cosses M8/M10", "1 lot", "10 000 – 20 000", "15 000"),
        ("P2 — Sécurité", "Disjoncteur DC + parafoudre", "Coffret DC panneaux (orage Niamey)", "1", "25 000 – 45 000", "35 000"),
        ("P3 — Option", "Coffret AC / différentiel 30 mA", "Protection sortie onduleur", "1", "20 000 – 35 000", "25 000"),
        ("P3 — Option", "Structure support panneaux", "Inclinaison ~15°, orientation Sud", "1", "40 000 – 80 000", "60 000"),
        ("TOTAL", "Budget d’appoint estimé (hors panneaux / onduleur déjà acquis)", "—", "—", "—", "~ 698 000"),
    ]
    for i, row_data in enumerate(achats):
        row = table.add_row()
        fill_row(row, row_data, zebra=(i % 2 == 1))
        if row_data[0] == "TOTAL":
            for cell in row.cells:
                shade_cell(cell, "FFF4EA")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = ORANGE
    add_para(
        doc,
        "Note : le budget d’appoint ci-dessus complète le matériel déjà en votre possession (panneaux + PWM + onduleur). "
        "Si vous optez pour une batterie Lithium 12 V 200 Ah au lieu du GEL 300 Ah, comptez environ +150 000 à +250 000 FCFA, "
        "mais une durée de vie et une tenue à la chaleur nettement meilleures.",
        size=9,
        color=GRAY,
        align="justify",
        space_before=6,
        space_after=10,
    )

    # 6. Config technique
    add_heading_bar(doc, "6. Configuration technique retenue")
    table = doc.add_table(rows=1, cols=2)
    fill_header_row(table.rows[0], ["Paramètre", "Valeur / choix"], fill="1B7A3D")
    config = [
        ("Tension système", "12 V DC"),
        ("Champ solaire", "1 080 Wc (3×250 Wc via MPPT + 1×330 Wc via PWM)"),
        ("Régulation", "MPPT 60 A (neuf) + PWM 30 A (existant) en parallèle"),
        ("Stockage recommandé", "Batterie GEL 12 V 300 Ah (DoD max 50 %)"),
        ("Conversion", "Onduleur 2 000 W 12 V → 230 V (Pur Sinus recommandé)"),
        ("Charge simultanée actuelle", "420–520 W (marge confortable)"),
        ("Autonomie cible", "≈ 3 heures (soirée)"),
        ("Évolution frigo", "Prévoir 2e batterie 200 Ah en parallèle ou passer Lithium"),
        ("Orientation / inclinaison", "Sud / ≈ 15° (Niamey)"),
    ]
    for i, row_data in enumerate(config):
        row = table.add_row()
        fill_row(row, row_data, zebra=(i % 2 == 1))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 7. Consignes
    add_heading_bar(doc, "7. Consignes d’installation & sécurité")
    consignes = [
        "Installer le fusible 125 A au plus près de la borne positive de la batterie (< 30 cm).",
        "Ne jamais dépasser ~500–600 W de charge continue avec un parc 12 V 300 Ah (risque de chute de tension).",
        "Vérifier que l’onduleur est bien en 12 V et de type Pur Sinus avant branchement du frigo ou Starlink.",
        "Ne pas mélanger les 4 panneaux sur un seul régulateur PWM : utiliser le schéma bi-régulateur ci-dessus.",
        "Protéger la batterie de la chaleur directe (local ventilé, hors soleil) — critique à Niamey.",
        "Serrer toutes les cosses ; vérifier le couple de serrage après 48 h de fonctionnement.",
    ]
    for c_text in consignes:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.3)
        run = p.add_run("• " + c_text)
        set_run_font(run, size=9, color=GRAY)

    # 8. Prochaines étapes
    add_para(doc, "", size=6, space_after=4)
    add_heading_bar(doc, "8. Prochaines étapes recommandées")
    etapes = [
        "Confirmer la tension (12 V / 24 V) et le type (Pur Sinus) de l’onduleur 2 000 W.",
        "Acheter en priorité : MPPT 60 A + batterie 300 Ah + câbles 25/35 mm² + fusible 125 A.",
        "Réaliser le câblage selon le schéma Figure 1, puis tests à vide puis en charge progressive.",
        "Planifier l’ajout frigo uniquement après renforcement batterie (2e 200 Ah ou Lithium).",
    ]
    for i, e in enumerate(etapes, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"{i}. {e}")
        set_run_font(run, size=10, color=NAVY)

    # Signature
    add_para(doc, "", size=8, space_after=8)
    add_para(doc, "Pour Niger Certify — Expertise technique", size=10, bold=True, color=NAVY, align="right")
    add_para(doc, "Document technique non contractuel — devis détaillé sur demande", size=8, color=GRAY, align="right")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"[OK] Document généré : {OUT}")


if __name__ == "__main__":
    build()
