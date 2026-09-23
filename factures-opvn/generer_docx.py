# -*- coding: utf-8 -*-
"""Génère la facture proforma OPVN en DOCX modifiable, dans deux styles.

    python3 generer_docx.py            # les deux styles
    python3 generer_docx.py moderne    # un seul style

Le rendu suit celui de generer_pdf.py : même contenu, même maquette.
"""

import os
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

import donnees_facture as D

SORTIE = os.environ.get("SORTIE_FACTURES", "/opt/cursor/artifacts")

LARGEUR_UTILE = 174  # mm

THEMES = {
    "moderne": {
        "libelle": "Moderne",
        "encre": "12263F",
        "accent": "0E9AA7",
        "doux": "EEF4F8",
        "alterne": "FBFCFD",
        "gris": "6B7A8C",
        "filet": "DCE4EC",
        "clair": "B9C4D0",
        "police_titre": "Calibri",
        "police_corps": "Calibri",
    },
    "elegante": {
        "libelle": "Élégante",
        "encre": "23272B",
        "accent": "A9762F",
        "doux": "FAF6EF",
        "alterne": "FDFBF7",
        "gris": "77706A",
        "filet": "E2DCD2",
        "clair": "C9C2BA",
        "police_titre": "Georgia",
        "police_corps": "Calibri",
    },
}


# --------------------------------------------------------------------------
# Utilitaires OOXML
#
# L'ordre des balises est imposé par le schéma WordprocessingML : un élément
# placé au mauvais rang rend le fichier illisible pour Word. Les constantes
# ci-dessous listent, en ordre de schéma, les balises qui doivent suivre
# celle que l'on insère.
# --------------------------------------------------------------------------

APRES_TC_BORDERS = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                    "w:tcFitText", "w:vAlign", "w:hideMark")
APRES_TC_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                "w:vAlign", "w:hideMark")
APRES_TC_MAR = ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
APRES_R_SPACING = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs",
                   "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
                   "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em",
                   "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
APRES_P_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
               "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
               "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
               "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
               "w:contextualSpacing", "w:mirrorIndents",
               "w:suppressOverlap", "w:jc", "w:textDirection",
               "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
               "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr")
APRES_P_TABS = APRES_P_BDR[2:]
APRES_TBL_W = ("w:tblJc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
               "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
               "w:tblCaption", "w:tblDescription")

COTES = ("top", "start", "bottom", "end")
SUIVANTS_BORDURE_PARAGRAPHE = {
    "top": ("w:left", "w:bottom", "w:right"),
    "left": ("w:bottom", "w:right"),
    "bottom": ("w:right",),
    "right": (),
}


def _inserer_avant(parent, element, successeurs):
    """Insère `element` avant le premier successeur présent (ordre du schéma)."""
    for tag in successeurs:
        cible = parent.find(qn(tag))
        if cible is not None:
            cible.addprevious(element)
            return element
    parent.append(element)
    return element


def _obtenir(parent, tag, successeurs):
    """Récupère l'enfant `tag`, ou l'insère à sa place dans le schéma."""
    element = parent.find(qn(tag))
    if element is None:
        element = _inserer_avant(parent, OxmlElement(tag), successeurs)
    return element


def _twips(millimetres):
    return str(int(round(millimetres * 56.7)))


def fond(cellule, couleur):
    shd = _obtenir(cellule._tc.get_or_add_tcPr(), "w:shd", APRES_TC_SHD)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)


def bordures_cellule(cellule, defaut_nil=False, **cotes):
    """cotes : top/start/bottom/end = (couleur, taille en 1/8 de point).

    Réécrit entièrement w:tcBorders : passer en un seul appel tous les côtés
    voulus.
    """
    tc_pr = cellule._tc.get_or_add_tcPr()
    ancien = tc_pr.find(qn("w:tcBorders"))
    if ancien is not None:
        tc_pr.remove(ancien)
    bordures = OxmlElement("w:tcBorders")
    for cote in COTES:
        if cote in cotes:
            couleur, taille = cotes[cote]
        elif defaut_nil:
            couleur, taille = "auto", 0
        else:
            continue
        element = OxmlElement(f"w:{cote}")
        element.set(qn("w:val"), "single" if taille else "nil")
        element.set(qn("w:sz"), str(taille))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), couleur)
        bordures.append(element)
    _inserer_avant(tc_pr, bordures, APRES_TC_BORDERS)


def marges_cellule(cellule, haut=3, bas=3, gauche=3, droite=3):
    """Marges internes de la cellule, en millimètres."""
    tc_pr = cellule._tc.get_or_add_tcPr()
    ancien = tc_pr.find(qn("w:tcMar"))
    if ancien is not None:
        tc_pr.remove(ancien)
    marges = OxmlElement("w:tcMar")
    for tag, valeur in (("top", haut), ("start", gauche),
                        ("bottom", bas), ("end", droite)):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:w"), _twips(valeur))
        element.set(qn("w:type"), "dxa")
        marges.append(element)
    _inserer_avant(tc_pr, marges, APRES_TC_MAR)


def largeurs_colonnes(table, largeurs_mm):
    """Fixe la grille, la largeur totale et la largeur de chaque cellule.

    À appeler une fois toutes les lignes ajoutées : une ligne créée après
    coup resterait à la largeur uniforme d'origine.
    """
    table.autofit = False
    total = _obtenir(table._tbl.tblPr, "w:tblW", APRES_TBL_W)
    total.set(qn("w:w"), _twips(sum(largeurs_mm)))
    total.set(qn("w:type"), "dxa")

    grille = table._tbl.find(qn("w:tblGrid"))
    if grille is not None:
        for colonne, largeur in zip(grille.findall(qn("w:gridCol")),
                                    largeurs_mm):
            colonne.set(qn("w:w"), _twips(largeur))

    for ligne in table.rows:
        for cellule, largeur in zip(ligne.cells, largeurs_mm):
            cellule.width = Mm(largeur)


def hauteur_ligne(ligne, millimetres):
    hauteur = OxmlElement("w:trHeight")
    hauteur.set(qn("w:val"), _twips(millimetres))
    hauteur.set(qn("w:hRule"), "atLeast")
    ligne._tr.get_or_add_trPr().append(hauteur)


def espacement_lettres(run, vingtiemes_de_point):
    spacing = _obtenir(run._element.get_or_add_rPr(), "w:spacing",
                       APRES_R_SPACING)
    spacing.set(qn("w:val"), str(vingtiemes_de_point))


def bordure_paragraphe(paragraphe, cote, couleur, taille, espace=4):
    bordures = _obtenir(paragraphe._p.get_or_add_pPr(), "w:pBdr", APRES_P_BDR)
    ancien = bordures.find(qn(f"w:{cote}"))
    if ancien is not None:
        bordures.remove(ancien)
    element = OxmlElement(f"w:{cote}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(taille))
    element.set(qn("w:space"), str(espace))
    element.set(qn("w:color"), couleur)
    _inserer_avant(bordures, element, SUIVANTS_BORDURE_PARAGRAPHE[cote])


def tabulation_droite(paragraphe, position_mm):
    tabs = _obtenir(paragraphe._p.get_or_add_pPr(), "w:tabs", APRES_P_TABS)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), _twips(position_mm))
    tabs.append(tab)


# --------------------------------------------------------------------------
# Écriture de texte
# --------------------------------------------------------------------------

def ecrire(paragraphe, texte, police="Calibri", taille=9, gras=False,
           couleur="000000", majuscules=False, interlettrage=None,
           italique=False):
    run = paragraphe.add_run(texte)
    run.font.name = police
    run.font.size = Pt(taille)
    run.font.bold = gras
    run.font.italic = italique
    run.font.color.rgb = RGBColor.from_string(couleur)
    run.font.all_caps = majuscules
    polices = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    if polices is not None:
        polices.set(qn("w:hAnsi"), police)
        polices.set(qn("w:cs"), police)
    if interlettrage:
        espacement_lettres(run, interlettrage)
    return run


def regler(paragraphe, avant=0, apres=0, interligne=1.0,
           alignement=WD_ALIGN_PARAGRAPH.LEFT):
    mise_en_forme = paragraphe.paragraph_format
    mise_en_forme.space_before = Pt(avant)
    mise_en_forme.space_after = Pt(apres)
    mise_en_forme.line_spacing = interligne
    paragraphe.alignment = alignement
    return paragraphe


def par_cellule(cellule, premier=True):
    """Renvoie un paragraphe de la cellule (réutilise le premier s'il est vide)."""
    if premier and cellule.paragraphs and not cellule.paragraphs[0].text:
        return cellule.paragraphs[0]
    return cellule.add_paragraph()


def paragraphe_espace(doc, points):
    """Paragraphe vide servant d'espacement vertical."""
    paragraphe = doc.add_paragraph()
    regler(paragraphe)
    paragraphe.paragraph_format.space_after = Pt(points)
    paragraphe.add_run("").font.size = Pt(1)
    return paragraphe


# --------------------------------------------------------------------------
# Blocs du document
# --------------------------------------------------------------------------

def entete_moderne(doc, t):
    table = doc.add_table(rows=1, cols=2)
    gauche, droite = table.rows[0].cells

    for cellule in (gauche, droite):
        fond(cellule, t["encre"])
        marges_cellule(cellule, haut=5, bas=5, gauche=5, droite=5)
        bordures_cellule(cellule, defaut_nil=True, bottom=(t["accent"], 30))
        cellule.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    paragraphe = par_cellule(gauche)
    regler(paragraphe, apres=2)
    ecrire(paragraphe, D.FOURNISSEUR["nom"], t["police_corps"], 15, True,
           "FFFFFF", interlettrage=4)
    for texte in (D.FOURNISSEUR["forme"],
                  f"{D.FOURNISSEUR['adresse']}  ·  "
                  f"Tél {D.FOURNISSEUR['tel']}",
                  f"NIF {D.FOURNISSEUR['nif']}"):
        paragraphe = par_cellule(gauche, premier=False)
        regler(paragraphe, apres=1)
        ecrire(paragraphe, texte, t["police_corps"], 7.5, False, t["clair"])

    paragraphe = par_cellule(droite)
    regler(paragraphe, alignement=WD_ALIGN_PARAGRAPH.RIGHT)
    ecrire(paragraphe, "PROFORMA", t["police_corps"], 8, True, t["accent"],
           interlettrage=30)
    paragraphe = par_cellule(droite, premier=False)
    regler(paragraphe, apres=1, alignement=WD_ALIGN_PARAGRAPH.RIGHT)
    ecrire(paragraphe, "FACTURE", t["police_corps"], 20, True, "FFFFFF",
           interlettrage=8)
    paragraphe = par_cellule(droite, premier=False)
    regler(paragraphe, alignement=WD_ALIGN_PARAGRAPH.RIGHT)
    ecrire(paragraphe, f"N° {D.FACTURE['numero']}", t["police_corps"], 8,
           False, t["clair"])

    largeurs_colonnes(table, [LARGEUR_UTILE * 0.62, LARGEUR_UTILE * 0.38])


def entete_elegante(doc, t):
    paragraphe = doc.add_paragraph()
    regler(paragraphe, apres=6, alignement=WD_ALIGN_PARAGRAPH.CENTER)
    bordure_paragraphe(paragraphe, "top", t["encre"], 12, espace=1)
    bordure_paragraphe(paragraphe, "bottom", t["encre"], 4, espace=6)
    ecrire(paragraphe, D.FOURNISSEUR["nom"], t["police_titre"], 14, True,
           t["encre"], interlettrage=12)

    for texte in (D.FOURNISSEUR["forme"],
                  f"{D.FOURNISSEUR['adresse']}  ·  "
                  f"Tél {D.FOURNISSEUR['tel']}  ·  "
                  f"NIF {D.FOURNISSEUR['nif']}"):
        paragraphe = doc.add_paragraph()
        regler(paragraphe, apres=1, alignement=WD_ALIGN_PARAGRAPH.CENTER)
        ecrire(paragraphe, texte, t["police_corps"], 7.5, False, t["gris"])

    paragraphe = doc.add_paragraph()
    regler(paragraphe, avant=6, apres=5,
           alignement=WD_ALIGN_PARAGRAPH.CENTER)
    bordure_paragraphe(paragraphe, "top", t["accent"], 6, espace=6)
    bordure_paragraphe(paragraphe, "bottom", t["accent"], 6, espace=6)
    ecrire(paragraphe, "FACTURE PROFORMA", t["police_titre"], 17, True,
           t["encre"], interlettrage=14)


def bloc_info(doc, t):
    items = [
        ("FACTURE N°", D.FACTURE["numero"]),
        ("DATE D'ÉMISSION", D.FACTURE["date"]),
        ("VALABLE JUSQU'AU", D.FACTURE["validite"]),
        ("DEVISE", "Franc CFA (FCFA)"),
    ]
    table = doc.add_table(rows=1, cols=4)
    for cellule, (libelle, valeur) in zip(table.rows[0].cells, items):
        fond(cellule, t["doux"])
        marges_cellule(cellule, haut=2.4, bas=2.4, gauche=3, droite=2)
        bordures_cellule(cellule, defaut_nil=True, top=(t["accent"], 18),
                         end=("FFFFFF", 8))
        paragraphe = par_cellule(cellule)
        regler(paragraphe, apres=2)
        ecrire(paragraphe, libelle, t["police_corps"], 6.5, True, t["gris"],
               interlettrage=12)
        paragraphe = par_cellule(cellule, premier=False)
        regler(paragraphe)
        ecrire(paragraphe, valeur, t["police_corps"], 9.5, True, t["encre"])
    largeurs_colonnes(table, [LARGEUR_UTILE / 4] * 4)


def bloc_client(doc, t):
    table = doc.add_table(rows=1, cols=2)
    paires = [("CLIENT", D.FACTURE["client"], t["accent"]),
              ("OBJET", D.FACTURE["objet"], t["filet"])]
    for cellule, (libelle, valeur, couleur_barre) in zip(
            table.rows[0].cells, paires):
        marges_cellule(cellule, haut=1, bas=1, gauche=3, droite=3)
        bordures_cellule(cellule, defaut_nil=True,
                         start=(couleur_barre, 18))
        paragraphe = par_cellule(cellule)
        regler(paragraphe, apres=2)
        ecrire(paragraphe, libelle, t["police_corps"], 6.5, True,
               t["accent"], interlettrage=12)
        paragraphe = par_cellule(cellule, premier=False)
        regler(paragraphe)
        ecrire(paragraphe, valeur, t["police_corps"], 9,
               libelle == "CLIENT", t["encre"])
    largeurs_colonnes(table, [LARGEUR_UTILE * 0.52, LARGEUR_UTILE * 0.48])


def titre_section(doc, t, texte):
    paragraphe = doc.add_paragraph()
    regler(paragraphe, apres=3)
    ecrire(paragraphe, texte, t["police_titre"], 10.5, True, t["encre"])


def tableau_lignes(doc, t):
    table = doc.add_table(rows=1, cols=5)

    entetes = ["N°", "DÉSIGNATION", "QTÉ", "PRIX UNIT.", "MONTANT"]
    for index, (cellule, libelle) in enumerate(
            zip(table.rows[0].cells, entetes)):
        fond(cellule, t["encre"])
        marges_cellule(cellule, haut=2, bas=2, gauche=2.2, droite=2.2)
        bordures_cellule(cellule, defaut_nil=True)
        cellule.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraphe = par_cellule(cellule)
        regler(paragraphe,
               alignement=WD_ALIGN_PARAGRAPH.RIGHT if index >= 2
               else WD_ALIGN_PARAGRAPH.LEFT)
        ecrire(paragraphe, libelle, t["police_corps"], 7, True, "FFFFFF",
               interlettrage=10)

    for index, (numero, titre, detail, qte, pu) in enumerate(D.LIGNES):
        cellules = table.add_row().cells
        for cellule in cellules:
            fond(cellule, t["alterne"] if index % 2 else "FFFFFF")
            marges_cellule(cellule, haut=2.2, bas=2.2, gauche=2.2,
                           droite=2.2)
            bordures_cellule(cellule, defaut_nil=True,
                             bottom=(t["filet"], 4))
            cellule.vertical_alignment = WD_ALIGN_VERTICAL.TOP

        paragraphe = par_cellule(cellules[0])
        regler(paragraphe, alignement=WD_ALIGN_PARAGRAPH.RIGHT)
        ecrire(paragraphe, f"{numero:02d}", t["police_corps"], 9, False,
               t["encre"])

        paragraphe = par_cellule(cellules[1])
        regler(paragraphe, apres=2)
        ecrire(paragraphe, titre, t["police_corps"], 9, True, t["encre"])
        paragraphe = par_cellule(cellules[1], premier=False)
        regler(paragraphe, interligne=1.05)
        ecrire(paragraphe, detail, t["police_corps"], 7.5, False, t["gris"])

        valeurs = [str(qte), D.montant(pu, devise=False),
                   D.montant(qte * pu, devise=False)]
        for cellule, valeur, gras in zip(cellules[2:], valeurs,
                                         (False, False, True)):
            paragraphe = par_cellule(cellule)
            regler(paragraphe, alignement=WD_ALIGN_PARAGRAPH.RIGHT)
            ecrire(paragraphe, valeur, t["police_corps"], 9, gras,
                   t["encre"])

    cellules = table.add_row().cells
    for cellule in cellules:
        fond(cellule, t["doux"])
        marges_cellule(cellule, haut=2.6, bas=2.6, gauche=2.2, droite=2.2)
        bordures_cellule(cellule, defaut_nil=True, top=(t["encre"], 12))
        cellule.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    paragraphe = par_cellule(cellules[1])
    regler(paragraphe)
    ecrire(paragraphe, "Total des articles", t["police_corps"], 9, True,
           t["encre"])
    for cellule, valeur in ((cellules[2], str(D.total_quantite())),
                            (cellules[4], D.montant(D.total(),
                                                    devise=False))):
        paragraphe = par_cellule(cellule)
        regler(paragraphe, alignement=WD_ALIGN_PARAGRAPH.RIGHT)
        ecrire(paragraphe, valeur, t["police_corps"], 9, True, t["encre"])

    largeurs_colonnes(table, [10, 94, 14, 24, 32])


def carte_total(doc, t):
    table = doc.add_table(rows=1, cols=2)
    gauche, droite = table.rows[0].cells
    for cellule in (gauche, droite):
        fond(cellule, t["encre"])
        marges_cellule(cellule, haut=4, bas=4, gauche=5, droite=5)
        bordures_cellule(cellule, defaut_nil=True)
        cellule.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    bordures_cellule(droite, defaut_nil=True, start=(t["clair"], 4))

    paragraphe = par_cellule(gauche)
    regler(paragraphe, apres=2)
    ecrire(paragraphe, "NET À PAYER", t["police_corps"], 6.5, True,
           t["clair"], interlettrage=14)
    paragraphe = par_cellule(gauche, premier=False)
    regler(paragraphe, apres=2)
    ecrire(paragraphe, D.montant(D.total()), t["police_corps"], 19, True,
           "FFFFFF")
    paragraphe = par_cellule(gauche, premier=False)
    regler(paragraphe)
    ecrire(paragraphe, D.TOTAL_EN_LETTRES, t["police_corps"], 7.5, False,
           t["clair"])

    for index, (libelle, valeur) in enumerate(D.CONDITIONS):
        paragraphe = par_cellule(droite, premier=(index == 0))
        regler(paragraphe, apres=3)
        ecrire(paragraphe, libelle, t["police_corps"], 7.5, True, "FFFFFF")
        ecrire(paragraphe, f" — {valeur}", t["police_corps"], 7.5, False,
               t["clair"])

    largeurs_colonnes(table, [LARGEUR_UTILE * 0.58, LARGEUR_UTILE * 0.42])


def bloc_note(doc, t):
    table = doc.add_table(rows=1, cols=1)
    cellule = table.rows[0].cells[0]
    fond(cellule, t["doux"])
    marges_cellule(cellule, haut=2.6, bas=2.6, gauche=3.6, droite=3.6)
    bordures_cellule(cellule, defaut_nil=True, start=(t["accent"], 18))
    paragraphe = par_cellule(cellule)
    regler(paragraphe, apres=2)
    ecrire(paragraphe, "NOTE", t["police_corps"], 6.5, True, t["accent"],
           interlettrage=12)
    paragraphe = par_cellule(cellule, premier=False)
    regler(paragraphe)
    ecrire(paragraphe, D.NOTE, t["police_corps"], 9, False, t["encre"])
    largeurs_colonnes(table, [LARGEUR_UTILE])


def bloc_signatures(doc, t):
    table = doc.add_table(rows=2, cols=2)
    hauteur_ligne(table.rows[0], 14)

    for index, cellule in enumerate(table.rows[0].cells):
        marges_cellule(cellule, haut=1, bas=1,
                       gauche=0 if index == 0 else 5,
                       droite=5 if index == 0 else 0)
        bordures_cellule(cellule, defaut_nil=True, bottom=(t["gris"], 6))
        regler(par_cellule(cellule))

    paires = [("LE FOURNISSEUR", f"E.I. {D.FOURNISSEUR['nom'].title()}"),
              ("LE CLIENT (OPVN)", "Lu et approuvé")]
    for index, (cellule, (libelle, valeur)) in enumerate(
            zip(table.rows[1].cells, paires)):
        marges_cellule(cellule, haut=1.6, bas=1,
                       gauche=0 if index == 0 else 5,
                       droite=5 if index == 0 else 0)
        bordures_cellule(cellule, defaut_nil=True)
        paragraphe = par_cellule(cellule)
        regler(paragraphe, apres=2)
        ecrire(paragraphe, libelle, t["police_corps"], 6.5, True, t["gris"],
               interlettrage=12)
        paragraphe = par_cellule(cellule, premier=False)
        regler(paragraphe)
        ecrire(paragraphe, valeur, t["police_corps"], 9, False, t["encre"])

    largeurs_colonnes(table, [LARGEUR_UTILE / 2] * 2)


def pied_de_page(doc, t):
    paragraphe = doc.sections[0].footer.paragraphs[0]
    regler(paragraphe)
    bordure_paragraphe(paragraphe, "top", t["filet"], 4, espace=6)
    tabulation_droite(paragraphe, LARGEUR_UTILE)
    ecrire(paragraphe,
           f"E.I. {D.FOURNISSEUR['nom'].title()}  ·  "
           f"NIF {D.FOURNISSEUR['nif']}",
           t["police_corps"], 7, False, t["gris"])
    ecrire(paragraphe, "\t", t["police_corps"], 7)
    ecrire(paragraphe,
           f"{D.FACTURE['numero']}  ·  {D.montant(D.total())}",
           t["police_corps"], 7, False, t["gris"])


# --------------------------------------------------------------------------
# Assemblage
# --------------------------------------------------------------------------

def construire(style):
    t = THEMES[style]
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = t["police_corps"]
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(16)
    section.footer_distance = Mm(9)

    proprietes = doc.core_properties
    proprietes.title = f"Facture proforma {D.FACTURE['numero']} — OPVN"
    proprietes.author = f"E.I. {D.FOURNISSEUR['nom'].title()}"
    proprietes.subject = "Confection de tenues avec logo OPVN"

    if style == "moderne":
        entete_moderne(doc, t)
        paragraphe_espace(doc, 10)
    else:
        entete_elegante(doc, t)
        paragraphe_espace(doc, 4)

    bloc_info(doc, t)
    paragraphe_espace(doc, 12)
    bloc_client(doc, t)
    paragraphe_espace(doc, 14)
    titre_section(doc, t, "Détail de la prestation")
    tableau_lignes(doc, t)
    paragraphe_espace(doc, 12)
    carte_total(doc, t)
    paragraphe_espace(doc, 10)
    bloc_note(doc, t)
    paragraphe_espace(doc, 8)
    bloc_signatures(doc, t)
    pied_de_page(doc, t)

    chemin = os.path.join(SORTIE,
                          f"Facture_Proforma_OPVN_style_{style}.docx")
    doc.save(chemin)
    return chemin


def main():
    os.makedirs(SORTIE, exist_ok=True)
    demandes = sys.argv[1:] or list(THEMES)
    for style in demandes:
        if style not in THEMES:
            raise SystemExit(
                f"Style inconnu : {style} (choix : {', '.join(THEMES)})")
        print("DOCX généré :", construire(style))


if __name__ == "__main__":
    main()
