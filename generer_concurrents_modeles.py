#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2 concurrentes avec modèles différents (A=vert Sahel, B=bordeaux NTP). Prix inchangés, base reste meilleure."""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def fmt(n):
    return f"{n:,}".replace(",", " ")

def bg(cell, hx):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hx)
    tcPr.append(shd)

def para(cell, text, size=9, bold=False, color=None, align=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p

DESC1 = "Ensemble T-shirt + Pantalon + Casquette avec logo OPVN\nPlanton : vert + noir + casquette noire | Manœuvre : vert + noir + casquette noire | Chauffeur : orange + noir + casquette noire | Gardiens : bleu marine + noir"
DESC2 = "Contre-veste marron avec logo OPVN\nPlanton (2ème complet) + Gardiens (2ème complet)"
DESC3 = "Bleu de travail mécanicien avec logo OPVN (les deux complets bleus)\nManœuvre + Mécaniciens"
DESC4 = "Jalabia kaki / Ensemble à poches avec logo OPVN\nChauffeurs (jalabia kaki) + Chauffeurs et Graisseurs"

# ============ MODELE A : bandeau vert ============
def generer_A(output):
    lignes = [(DESC1, 200, 33000), (DESC2, 108, 38250), (DESC3, 24, 22750), (DESC4, 80, 30500)]
    total = sum(q*pu for _, q, pu in lignes)
    numero, dj, dv = "PF-2026-0914-002", "14/09/2026", "14/10/2026"
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Mm(10); sec.bottom_margin = Mm(12); sec.left_margin = Mm(14); sec.right_margin = Mm(14)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)

    # Bandeau plein largeur vert
    band = doc.add_table(rows=1, cols=1); band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.rows[0].cells[0]
    bg(c, "1B7A3D")
    c.text = ""
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ETS SAHEL CONFECTION & SERVICES"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Confection — Broderie — Tenues professionnelles  •  Quartier Lazaret, Niamey  •  Tél : +227 90 12 34 56\nsahel.confection@ne.com  —  NIF : 45210/S")
    r2.font.size = Pt(8); r2.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

    # Ligne info proforma décalée à droite + objet à gauche (style différent : 2 blocs séparés)
    doc.add_paragraph("")
    info = doc.add_table(rows=1, cols=2); info.alignment = WD_TABLE_ALIGNMENT.CENTER
    para(info.rows[0].cells[0], "Devis proforma — Client : OPVN (Office des Produits Vivriers du Niger, Niamey)\nObjet : Tenues de travail avec logo OPVN — Plantons, Manœuvres, Chauffeurs, Gardiens, Mécaniciens, Graisseurs", 8.5, False, (0x33,0x41,0x55), WD_ALIGN_PARAGRAPH.LEFT)
    bg(info.rows[0].cells[0], "EAF5EC")
    para(info.rows[0].cells[1], f"N° {numero}\nDate : {dj}\nÉchéance : {dv}\nValidité : 30 jours", 9, True, (0x1B,0x7A,0x3D), WD_ALIGN_PARAGRAPH.RIGHT)
    bg(info.rows[0].cells[1], "FFFFFF")

    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("Facture proforma en francs CFA"); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x1B,0x7A,0x3D)
    # filet vert
    from docx.shared import Pt as Pt2
    # tableau articles : en-tête vert
    tab = doc.add_table(rows=1, cols=5); tab.style = "Table Grid"; tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["RÉF", "DESCRIPTION", "QTÉ", "P.U. FCFA", "TOTAL FCFA"]):
        cc = tab.rows[0].cells[i]
        para(cc, h, 8.5, True, (0xFF,0xFF,0xFF), WD_ALIGN_PARAGRAPH.CENTER); bg(cc, "1B7A3D")
    for idx, (des, qte, pu) in enumerate(lignes, start=1):
        row = tab.add_row().cells
        parts = des.split("\n")
        para(row[0], f"A-{idx:02d}", 8.5, True, None, WD_ALIGN_PARAGRAPH.CENTER)
        p1 = para(row[1], parts[0], 8.5, True); p1.add_run("\n" + parts[1]).font.size = Pt(7.5)
        # couleur texte secondaire via runs
        for run in p1.runs[1:]:
            run.font.color.rgb = RGBColor(0x4B,0x5B,0x6B)
        para(row[2], str(qte), 8.5, False, None, WD_ALIGN_PARAGRAPH.CENTER)
        para(row[3], fmt(pu), 8.5, False, None, WD_ALIGN_PARAGRAPH.RIGHT)
        para(row[4], fmt(qte*pu), 8.5, True, (0x1B,0x7A,0x3D), WD_ALIGN_PARAGRAPH.RIGHT)
        if idx % 2 == 0:
            for cc in row: bg(cc, "F2F9F3")
    doc.add_paragraph("")
    tt = doc.add_table(rows=2, cols=2); tt.style = "Table Grid"; tt.alignment = WD_TABLE_ALIGNMENT.RIGHT
    para(tt.rows[0].cells[0], "Montant total HT / Net à payer (FCFA)", 9, True, None, WD_ALIGN_PARAGRAPH.RIGHT)
    para(tt.rows[0].cells[1], fmt(total), 10, True, (0x1B,0x7A,0x3D), WD_ALIGN_PARAGRAPH.RIGHT)
    bg(tt.rows[0].cells[0], "EAF5EC"); bg(tt.rows[0].cells[1], "EAF5EC")
    para(tt.rows[1].cells[0], "TVA", 8.5, False, None, WD_ALIGN_PARAGRAPH.RIGHT)
    para(tt.rows[1].cells[1], "Non applicable", 8.5, False, (0x66,0x70,0x85), WD_ALIGN_PARAGRAPH.RIGHT)
    p = doc.add_paragraph()
    rr = p.add_run(f"Arrêté à : Treize millions sept cent dix-sept mille ({fmt(total)}) francs CFA.")
    rr.bold = True; rr.font.size = Pt(8.5)
    doc.add_paragraph("Répartition : Planton vert/noir/casquette + contre-veste marron • Manœuvre vert/noir/casquette + bleu • Chauffeur orange/noir/casquette + jalabia kaki/poches • Gardiens bleu marine/noir + contre-veste marron • Mécaniciens bleu/bleu • Tous avec logo OPVN. Paiement : à convenir. Délai : à convenir.", style="List Bullet")
    doc.add_paragraph("")
    sig = doc.add_table(rows=1, cols=2); sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    para(sig.rows[0].cells[0], "Pour Sahel Confection\n\n\nCachet & Signature", 9, True, (0x1B,0x7A,0x3D), WD_ALIGN_PARAGRAPH.CENTER)
    para(sig.rows[0].cells[1], "Pour OPVN — Bon pour accord\n\n\nCachet & Signature", 9, False, None, WD_ALIGN_PARAGRAPH.CENTER)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"Ets Sahel Confection — {numero} — {fmt(total)} FCFA — Tél +227 90 12 34 56")
    r.font.size = Pt(7); r.font.color.rgb = RGBColor(0x1B,0x7A,0x3D)
    doc.save(output); print(f"OK A -> {output} {total}")
    return total

# ============ MODELE B : style bordeaux / carte ============
def generer_B(output):
    lignes = [(DESC1, 200, 32000), (DESC2, 108, 38500), (DESC3, 24, 24500), (DESC4, 80, 30500)]
    total = sum(q*pu for _, q, pu in lignes)
    numero, dj, dv = "PF-2026-0915-007", "15/09/2026", "15/10/2026"
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Mm(12); sec.bottom_margin = Mm(12); sec.left_margin = Mm(16); sec.right_margin = Mm(16)
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(9)

    # En-tête carte : logo + société
    head = doc.add_table(rows=1, cols=2); head.alignment = WD_TABLE_ALIGNMENT.CENTER
    logo = head.rows[0].cells[0]
    bg(logo, "7A1C1C")
    logo.text = ""
    pl = logo.paragraphs[0]; pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl = pl.add_run("NTP"); rl.bold = True; rl.font.size = Pt(26); rl.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    pl2 = logo.add_paragraph(); pl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl2 = pl2.add_run("NIGER TEXTILE PRO"); rl2.bold = True; rl2.font.size = Pt(9); rl2.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    co = head.rows[0].cells[1]
    co.text = ""
    pc = co.paragraphs[0]; pc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rc = pc.add_run("Niger Textile Pro – NTP\n"); rc.bold = True; rc.font.size = Pt(12); rc.font.color.rgb = RGBColor(0x7A,0x1C,0x1C)
    pc2 = co.add_paragraph()
    rc2 = pc2.add_run("Habillement professionnel & Broderie\nAvenue de l'Indépendance, Niamey — Tél : +227 96 78 45 12\ncontact@ntp-niger.ne — NIF : 38902/R — RC : 1234")
    rc2.font.size = Pt(8); rc2.font.color.rgb = RGBColor(0x44,0x44,0x44)

    # Titre encadré différent : centré avec filets
    doc.add_paragraph("")
    tt1 = doc.add_paragraph(); tt1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tt1.add_run("—  FACTURE PROFORMA  —"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x7A,0x1C,0x1C)
    tt2 = doc.add_paragraph(); tt2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tt2.add_run(f"N° {numero}  |  Émise le {dj}  |  Valable jusqu'au {dv}  |  Montants en FCFA")
    r.font.size = Pt(8); r.italic = True; r.font.color.rgb = RGBColor(0x66,0x70,0x85)
    tt3 = doc.add_paragraph(); tt3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tt3.add_run("Client : OPVN — Office des Produits Vivriers du Niger  •  Objet : confection de tenues avec logo OPVN")
    r.bold = True; r.font.size = Pt(8.5)

    tab = doc.add_table(rows=1, cols=4); tab.style = "Table Grid"; tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Désignation", "Qté", "Prix unit. (FCFA)", "Montant (FCFA)"]):
        cc = tab.rows[0].cells[i]
        para(cc, h, 8.5, True, (0xFF,0xFF,0xFF), WD_ALIGN_PARAGRAPH.CENTER); bg(cc, "7A1C1C")
    for idx, (des, qte, pu) in enumerate(lignes, start=1):
        row = tab.add_row().cells
        parts = des.split("\n")
        p0 = para(row[0], f"{idx}. {parts[0]}", 8, True, (0x1A,0x1A,0x1A))
        p0.add_run("\n" + parts[1]).font.size = Pt(7.5)
        para(row[1], str(qte), 9, False, None, WD_ALIGN_PARAGRAPH.CENTER)
        para(row[2], fmt(pu), 9, False, None, WD_ALIGN_PARAGRAPH.RIGHT)
        para(row[3], fmt(qte*pu), 9, True, (0x7A,0x1C,0x1C), WD_ALIGN_PARAGRAPH.RIGHT)
        if idx % 2 == 1:
            for cc in row: bg(cc, "FDF5F5")
    doc.add_paragraph("")
    tt = doc.add_table(rows=1, cols=2); tt.style = "Table Grid"; tt.alignment = WD_TABLE_ALIGNMENT.CENTER
    para(tt.rows[0].cells[0], f"TOTAL NET À PAYER : {fmt(total)} FCFA\nTreize millions cinq cent quatre-vingt-six mille francs CFA", 9, True, (0xFF,0xFF,0xFF), WD_ALIGN_PARAGRAPH.CENTER)
    bg(tt.rows[0].cells[0], "7A1C1C")
    para(tt.rows[0].cells[1], "TVA : —\nPaiement : espèces / virement / chèque\nDélai de confection : à convenir", 8, False, None, WD_ALIGN_PARAGRAPH.LEFT)
    p = doc.add_paragraph()
    r = p.add_run("Note : ")
    r.bold = True; r.font.size = Pt(8.5)
    rr = p.add_run("plantons, manœuvres, chauffeurs, gardiens, mécaniciens et graisseurs — tous les articles brodés avec le logo OPVN.")
    rr.font.size = Pt(8.5)
    doc.add_paragraph("")
    sig = doc.add_table(rows=1, cols=2); sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    para(sig.rows[0].cells[0], "Signature Fournisseur (NTP)\n\n\n", 9, False, None, WD_ALIGN_PARAGRAPH.CENTER)
    para(sig.rows[0].cells[1], "Signature Client (OPVN)\nLu et approuvé\n\n", 9, False, None, WD_ALIGN_PARAGRAPH.CENTER)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"Niger Textile Pro — {numero} — {fmt(total)} FCFA")
    r.font.size = Pt(7); r.font.color.rgb = RGBColor(0x7A,0x1C,0x1C)
    doc.save(output); print(f"OK B -> {output} {total}")
    return total

if __name__ == "__main__":
    a = generer_A("/workspace/Facture_Proforma_OPVN_Concurrent_A.docx")
    b = generer_B("/workspace/Facture_Proforma_OPVN_Concurrent_B.docx")
    print(f"Base=12420000 A={a} B={b}")
