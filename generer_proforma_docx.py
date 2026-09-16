#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génération facture proforma OPVN en DOCX (Word)"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "/workspace/Facture_Proforma_OPVN_FCFA.docx"
NUMERO = "PF-2026-0916-001"
DATE_JOUR = "16/09/2026"
DATE_VAL = "16/10/2026"

def fmt(n):
    return f"{n:,}".replace(",", " ")

def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for m, v in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{m}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)

lignes = [
    ("Ensemble T-shirt + Pantalon + Casquette avec logo OPVN\nPlanton : T-shirt vert + pantalon noir + casquette noire | Manœuvre : T-shirt vert + pantalon noir + casquette noire | Chauffeur (1er complet) : T-shirt orange + pantalon noir + casquette noire | Gardiens (1er complet) : T-shirt bleu marine + pantalon noir", 200, 30000),
    ("Contre-veste marron avec logo OPVN\nPlanton (2ème complet) + Gardiens (2ème complet)", 108, 35000),
    ("Bleu de travail mécanicien avec logo OPVN (les deux complets bleus)\nManœuvre : bleu comme mécanicien + Mécaniciens : bleu / bleu", 24, 20000),
    ("Jalabia kaki / Ensemble à poches avec logo OPVN\nChauffeurs (2ème complet : jalabia kaki) + Chauffeurs et Graisseurs : ensemble à poches", 80, 27000),
]
total = sum(q*pu for _, q, pu in lignes)

doc = Document()
for section in doc.sections:
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(12)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(9)

# En-tête émetteur + bloc proforma
t_head = doc.add_table(rows=1, cols=2)
t_head.alignment = WD_TABLE_ALIGNMENT.CENTER
c0, c1 = t_head.rows[0].cells
c0.text = "[Nom de votre entreprise]\nConfection & Fourniture de tenues professionnelles\nAdresse : [Quartier, Ville] — Tél : [+227 XX XX XX XX]\nE-mail : [contact@entreprise.ne] — NIF : [XXXXX]"
c1.text = f"FACTURE PROFORMA\nN° : {NUMERO}\nDate : {DATE_JOUR}\nValidité : {DATE_VAL} (30 jours)"
for p in c0.paragraphs:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
for p in c1.paragraphs:
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_cell_bg(c1, "E8F0FE")

doc.add_paragraph("")
titre = doc.add_paragraph()
titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = titre.add_run("FACTURE PROFORMA — TENUES DE TRAVAIL AVEC LOGO OPVN")
r.bold = True
r.font.size = Pt(15)
r.font.color.rgb = RGBColor(0x0F, 0x2A, 0x44)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = sub.add_run("Établie en Francs CFA (FCFA) — Tous les articles avec logo OPVN")
rr.font.size = Pt(9)
rr.font.color.rgb = RGBColor(0x5A, 0x6C, 0x7D)

# Client / Objet
t_info = doc.add_table(rows=1, cols=2)
t_info.style = "Table Grid"
t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
a, b = t_info.rows[0].cells
a.text = "Client : OPVN\nOffice des Produits Vivriers du Niger\nNiamey — Niger"
b.text = "Objet : Confection et fourniture de tenues de travail\nPlantons, Manœuvres, Chauffeurs, Gardiens,\nMécaniciens et Graisseurs — avec logo OPVN"
set_cell_bg(a, "F2F4F7")
set_cell_bg(b, "F2F4F7")

doc.add_paragraph("")
# Tableau articles
t = doc.add_table(rows=1, cols=5)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ["N°", "DÉSIGNATION", "QTÉ", "P.U. (FCFA)", "MONTANT (FCFA)"]
for i, h in enumerate(hdr):
    cell = t.rows[0].cells[i]
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_cell_bg(cell, "0F2A44")

widths = [Mm(10), Mm(90), Mm(15), Mm(30), Mm(35)]
for idx, (des, qte, pu) in enumerate(lignes, start=1):
    row = t.add_row().cells
    row[0].text = ""
    p0 = row[0].paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER; p0.add_run(str(idx)).font.size = Pt(8.5)
    row[1].text = ""
    p1 = row[1].paragraphs[0]
    parts = des.split("\n")
    run1 = p1.add_run(parts[0]); run1.bold = True; run1.font.size = Pt(8.5)
    if len(parts) > 1:
        run2 = p1.add_run("\n" + parts[1]); run2.font.size = Pt(7.5); run2.font.color.rgb = RGBColor(0x34, 0x40, 0x54)
    row[2].text = ""
    p2 = row[2].paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; p2.add_run(str(qte)).font.size = Pt(8.5)
    row[3].text = ""
    p3 = row[3].paragraphs[0]; p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p3.add_run(fmt(pu)).font.size = Pt(8.5)
    row[4].text = ""
    p4 = row[4].paragraphs[0]; p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p4.add_run(fmt(qte*pu)).font.size = Pt(8.5)
    if idx % 2 == 0:
        for c in row:
            set_cell_bg(c, "F8FAFC")

# Totaux
doc.add_paragraph("")
t_tot = doc.add_table(rows=3, cols=2)
t_tot.style = "Table Grid"
t_tot.alignment = WD_TABLE_ALIGNMENT.RIGHT
t_tot.rows[0].cells[0].text = "Total HT (FCFA)"
t_tot.rows[0].cells[1].text = fmt(total)
t_tot.rows[1].cells[0].text = "TVA"
t_tot.rows[1].cells[1].text = "—"
t_tot.rows[2].cells[0].text = "Total TTC / Net à payer (FCFA)"
t_tot.rows[2].cells[1].text = fmt(total)
for row in t_tot.rows:
    row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_cell_bg(t_tot.rows[2].cells[0], "E8F0FE")
set_cell_bg(t_tot.rows[2].cells[1], "E8F0FE")
for p in t_tot.rows[2].cells[0].paragraphs[0].runs:
    p.bold = True
for p in t_tot.rows[2].cells[1].paragraphs[0].runs:
    p.bold = True

p = doc.add_paragraph()
p.add_run("Arrêté la présente facture proforma à la somme de : ").font.size = Pt(8.5)
r = p.add_run("Douze millions quatre cent vingt mille (12 420 000) francs CFA.")
r.bold = True
r.font.size = Pt(8.5)

p = doc.add_paragraph()
r = p.add_run("Détail des répartitions :")
r.bold = True
r.font.size = Pt(9)
doc.add_paragraph("• Planton : T-shirt vert + pantalon noir + casquette noire, + contre-veste marron (2ème complet).\n• Manœuvre : T-shirt vert + pantalon noir + casquette noire, + bleu comme mécaniciens.\n• Chauffeur : T-shirt orange + pantalon noir + casquette noire, + jalabia kaki / ensemble à poches (2ème complet).\n• Gardiens : T-shirt bleu marine + pantalon noir, + contre-veste marron (2ème complet).\n• Mécaniciens : bleu / bleu (les deux complets).\n• NB : Tous les articles avec le logo de l’OPVN.", style="List Bullet").style.font.size = Pt(8.5)

p = doc.add_paragraph()
r = p.add_run("Conditions : ")
r.bold = True
r.font.size = Pt(8.5)
p.add_run("Prix en FCFA. Validité de l’offre : 30 jours. Délai de confection à convenir. Paiement : à préciser (espèces / virement / chèque).").font.size = Pt(8.5)

doc.add_paragraph("")
t_sig = doc.add_table(rows=1, cols=2)
t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
t_sig.rows[0].cells[0].text = "Le Fournisseur\n\n\nSignature & Cachet"
t_sig.rows[0].cells[1].text = "Le Client — OPVN\nLu et approuvé, Bon pour accord\n\n\nSignature & Cachet"
for c in t_sig.rows[0].cells:
    for par in c.paragraphs:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Document établi à titre proforma — ne tient pas lieu de facture définitive. Merci de nous retourner un exemplaire signé.")
r.font.size = Pt(7.5)
r.font.color.rgb = RGBColor(0x66, 0x70, 0x85)

# Pied de page
section = doc.sections[0]
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(f"Proforma N° {NUMERO} — {fmt(total)} FCFA — OPVN — Tenues avec logo")
r.font.size = Pt(7)
r.font.color.rgb = RGBColor(0x66, 0x70, 0x85)

doc.save(OUTPUT)
print(f"OK -> {OUTPUT} total={total}")
